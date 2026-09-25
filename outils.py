import smtplib
from email.mime.text import MIMEText
import requests
import os
from datetime import datetime
from tavily import TavilyClient
from groq import Groq
tavily_client = TavilyClient(api_key=os.getenv("TAVILY_API_KEY"))
def lire_fichier(chemin: str) -> str:
    from datetime import datetime
    try:
        with open(chemin, "r", encoding="utf-8") as f:
            return f.read()
    except Exception as e:
        return f"Erreur : {e}"

def ecrire_fichier(chemin: str, contenu: str) -> str:
    try:
        with open(chemin, "w", encoding="utf-8") as f:
            f.write(contenu)
        return f"Fichier {chemin} écrit avec succès."
    except Exception as e:
        return f"Erreur : {e}"

def lister_fichiers(dossier: str = ".") -> str:
    try:
        return "\n".join(os.listdir(dossier))
    except Exception as e:
        return f"Erreur : {e}"
def calculer(expression: str) -> str:
    try:
        resultat = eval(expression, {"__builtins__": {}})
        return str(resultat)
    except Exception as e:
        return f"Erreur de calcul : {e}"
def ajouter_a_memoire(information: str, categorie: str = "general") -> str:
    try:
        ligne = f"[{categorie}] {information}"
        
        # Vérifier si l'info existe déjà (évite les doublons)
        try:
            with open("memoire.txt", "r", encoding="utf-8") as f:
                contenu_existant = f.read()
            if information.strip() in contenu_existant:
                return "Cette information est déjà mémorisée."
        except FileNotFoundError:
            pass
        
        with open("memoire.txt", "a", encoding="utf-8") as f:
            f.write(ligne + "\n")
        return "Information mémorisée avec succès."
    except Exception as e:
        return f"Erreur : {e}"
def resumer_memoire() -> str:
    try:
        with open("memoire.txt", "r", encoding="utf-8") as f:
            contenu = f.read()
        lignes = [l for l in contenu.split("\n") if l.strip()]
        if len(lignes) < 20:
            return f"La mémoire contient {len(lignes)} entrées, pas besoin de résumé pour l'instant."
        return f"La mémoire contient {len(lignes)} entrées. Contenu complet :\n{contenu}"
    except FileNotFoundError:
        return "Aucune mémoire enregistrée pour l'instant."
def condenser_memoire_si_necessaire(client, seuil=30):
    try:
        with open("memoire.txt", "r", encoding="utf-8") as f:
            lignes = [l for l in f.read().split("\n") if l.strip()]
        
        if len(lignes) < seuil:
            return
        
        contenu = "\n".join(lignes)
        reponse = client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{
                "role": "user",
                "content": f"Condense ces informations mémorisées sur un utilisateur en gardant les catégories [etudes], [preferences], [taches], [general], en éliminant les doublons et redondances, réponse uniquement le texte condensé sans commentaire :\n\n{contenu}"
            }],
            max_tokens=1024
        )
        
        resume = reponse.choices[0].message.content
        with open("memoire.txt", "w", encoding="utf-8") as f:
            f.write(resume)
    except Exception:
        pass
def rechercher_web(requete: str) -> str:
    try:
        resultats = tavily_client.search(requete, max_results=3)
        return str(resultats)
    except Exception as e:
        return f"Erreur : {e}"
def obtenir_date_heure() -> str:
    return datetime.now().strftime("%A %d %B %Y, %H:%M")
def obtenir_date_heure_fuseau(fuseau: str = "Africa/Porto-Novo") -> str:
    try:
        from zoneinfo import ZoneInfo
        maintenant = datetime.now(ZoneInfo(fuseau))
        return maintenant.strftime("%A %d %B %Y, %H:%M") + f" ({fuseau})"
    except Exception as e:
        return f"Erreur : fuseau horaire invalide ou inconnu ({fuseau}). Utilise un format comme 'Europe/Paris', 'America/New_York', 'Asia/Tokyo'. Détail : {e}"
def ajouter_evenement(titre: str, date_heure: str, description: str = "") -> str:
    try:
        ligne = f"{date_heure} | {titre} | {description}\n"
        with open("agenda.txt", "a", encoding="utf-8") as f:
            f.write(ligne)
        return f"Événement '{titre}' ajouté pour le {date_heure}."
    except Exception as e:
        return f"Erreur : {e}"

def lister_evenements() -> str:
    try:
        with open("agenda.txt", "r", encoding="utf-8") as f:
            contenu = f.read().strip()
        if not contenu:
            return "Aucun événement dans l'agenda."
        return contenu
    except FileNotFoundError:
        return "Aucun événement dans l'agenda."

def supprimer_evenement(titre: str) -> str:
    try:
        with open("agenda.txt", "r", encoding="utf-8") as f:
            lignes = f.readlines()
        lignes_restantes = [l for l in lignes if titre.lower() not in l.lower()]
        if len(lignes_restantes) == len(lignes):
            return f"Aucun événement trouvé avec le titre '{titre}'."
        with open("agenda.txt", "w", encoding="utf-8") as f:
            f.writelines(lignes_restantes)
        return f"Événement(s) contenant '{titre}' supprimé(s)."
    except FileNotFoundError:
        return "Aucun agenda existant."
    except Exception as e:
        return f"Erreur : {e}"
def creer_flashcard(question: str, reponse: str) -> str:
    try:
        with open("flashcards.txt", "a", encoding="utf-8") as f:
            f.write(f"Q: {question}\nR: {reponse}\n---\n")
        return "Flashcard ajoutée."
    except Exception as e:
        return f"Erreur : {e}"
def convertir_unite(valeur: float, de_unite: str, vers_unite: str) -> str:
    conversions = {
        ("km", "m"): 1000, ("m", "km"): 0.001,
        ("m", "cm"): 100, ("cm", "m"): 0.01,
        ("kg", "g"): 1000, ("g", "kg"): 0.001,
        ("l", "ml"): 1000, ("ml", "l"): 0.001,
        ("h", "min"): 60, ("min", "h"): 1/60,
        ("min", "s"): 60, ("s", "min"): 1/60,
        ("km", "mile"): 0.621371, ("mile", "km"): 1.60934,
        ("kg", "lb"): 2.20462, ("lb", "kg"): 0.453592,
    }
    if de_unite == "celsius" and vers_unite == "fahrenheit":
        return str(valeur * 9/5 + 32)
    if de_unite == "fahrenheit" and vers_unite == "celsius":
        return str((valeur - 32) * 5/9)
    facteur = conversions.get((de_unite, vers_unite))
    if facteur:
        return str(valeur * facteur)
    return "Conversion non supportée"
def compter_mots(texte: str) -> str:
    nb_mots = len(texte.split())
    nb_caracteres = len(texte)
    return f"{nb_mots} mots, {nb_caracteres} caractères"
def obtenir_meteo(ville: str) -> str:
    try:
        cle = os.getenv("OPENWEATHER_API_KEY")
        url = f"https://api.openweathermap.org/data/2.5/weather?q={ville}&appid={cle}&units=metric&lang=fr"
        reponse = requests.get(url)
        data = reponse.json()
        if reponse.status_code != 200:
            return f"Erreur : {data.get('message', 'ville introuvable')}"
        temp = data["main"]["temp"]
        description = data["weather"][0]["description"]
        return f"À {ville} : {temp}°C, {description}"
    except Exception as e:
        return f"Erreur : {e}"
def envoyer_email(destinataire: str, sujet: str, message: str) -> str:
    try:
        expediteur = os.getenv("EMAIL_ADRESSE")
        mot_de_passe = os.getenv("EMAIL_MOT_DE_PASSE")

        destinataires = [d.strip() for d in destinataire.split(",")]

        mail = MIMEText(message)
        mail["Subject"] = sujet
        mail["From"] = expediteur
        mail["To"] = ", ".join(destinataires)

        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as serveur:
            serveur.login(expediteur, mot_de_passe)
            serveur.sendmail(expediteur, destinataires, mail.as_string())

        return f"Email envoyé à {', '.join(destinataires)}."
    except Exception as e:
        return f"Erreur : {e}"
def creer_skill(nom: str, contenu: str, description: str) -> str:
    try:
        import os
        if not os.path.exists("skills"):
            os.makedirs("skills")
        with open(f"skills/{nom}.txt", "w", encoding="utf-8") as f:
            f.write(f"DESCRIPTION: {description}\n{contenu}")
        return f"Skill '{nom}' créé avec succès."
    except Exception as e:
        return f"Erreur : {e}"
def lire_pdf(chemin: str) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(chemin)
        texte = ""
        for page in reader.pages:
            texte += page.extract_text() + "\n"
        return texte[:5000] if len(texte) > 5000 else texte
    except Exception as e:
        return f"Erreur : {e}"
def lire_docx(chemin: str) -> str:
    try:
        from docx import Document
        doc = Document(chemin)
        texte = "\n".join(p.text for p in doc.paragraphs)
        return texte[:5000] if len(texte) > 5000 else texte
    except Exception as e:
        return f"Erreur : {e}"

def ecrire_docx(chemin: str, contenu: str, titre: str = "") -> str:
    try:
        from docx import Document
        doc = Document()
        if titre:
            doc.add_heading(titre, level=1)
        for paragraphe in contenu.split("\n"):
            doc.add_paragraph(paragraphe)
        doc.save(chemin)
        return f"Document Word {chemin} créé avec succès."
    except Exception as e:
        return f"Erreur : {e}"

def lire_excel(chemin: str) -> str:
    try:
        from openpyxl import load_workbook
        classeur = load_workbook(chemin, data_only=True)
        resultat = []
        for feuille in classeur.sheetnames:
            ws = classeur[feuille]
            resultat.append(f"--- Feuille : {feuille} ---")
            for ligne in ws.iter_rows(values_only=True, max_row=50):
                resultat.append(" | ".join(str(c) if c is not None else "" for c in ligne))
        texte = "\n".join(resultat)
        return texte[:5000] if len(texte) > 5000 else texte
    except Exception as e:
        return f"Erreur : {e}"

def ecrire_excel(chemin: str, donnees: str, feuille: str = "Feuille1") -> str:
    try:
        from openpyxl import Workbook
        classeur = Workbook()
        ws = classeur.active
        ws.title = feuille
        for i, ligne in enumerate(donnees.split("\n"), start=1):
            for j, valeur in enumerate(ligne.split(","), start=1):
                ws.cell(row=i, column=j, value=valeur.strip())
        classeur.save(chemin)
        return f"Fichier Excel {chemin} créé avec succès."
    except Exception as e:
        return f"Erreur : {e}"

def lire_pptx(chemin: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(chemin)
        resultat = []
        for i, diapo in enumerate(prs.slides, start=1):
            resultat.append(f"--- Diapositive {i} ---")
            for forme in diapo.shapes:
                if forme.has_text_frame:
                    resultat.append(forme.text_frame.text)
        texte = "\n".join(resultat)
        return texte[:5000] if len(texte) > 5000 else texte
    except Exception as e:
        return f"Erreur : {e}"

def creer_pptx(chemin: str, titre: str, contenu: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation()
        diapo = prs.slides.add_slide(prs.slide_layouts[1])
        diapo.shapes.title.text = titre
        zone_texte = diapo.placeholders[1]
        zone_texte.text = contenu
        prs.save(chemin)
        return f"Présentation PowerPoint {chemin} créée avec succès."
    except Exception as e:
        return f"Erreur : {e}"
def suivre_progression(matiere: str, statut: str) -> str:
    try:
        ligne = f"[{matiere}] {statut} - {datetime.now().strftime('%Y-%m-%d')}\n"
        with open("progression.txt", "a", encoding="utf-8") as f:
            f.write(ligne)
        return "Progression enregistrée."
    except Exception as e:
        return f"Erreur : {e}"

def voir_progression() -> str:
    try:
        with open("progression.txt", "r", encoding="utf-8") as f:
            return f.read()
    except FileNotFoundError:
        return "Aucune progression enregistrée."
def exporter_donnees() -> str:
    try:
        import json
        donnees = {}
        
        for nom_fichier in ["memoire.txt", "progression.txt", "flashcards.txt"]:
            try:
                with open(nom_fichier, "r", encoding="utf-8") as f:
                    donnees[nom_fichier] = f.read()
            except FileNotFoundError:
                donnees[nom_fichier] = ""
        
        with open("export_complet.json", "w", encoding="utf-8") as f:
            json.dump(donnees, f, ensure_ascii=False, indent=2)
        
        return "Export créé : export_complet.json. Utilise lire_fichier pour le consulter, ou demande-moi de te l'envoyer par email."
    except Exception as e:
        return f"Erreur : {e}"
def planifier_revisions(sujets: str, date_examen: str) -> str:
    try:
        from datetime import datetime
        
        liste_sujets = [s.strip() for s in sujets.split(",") if s.strip()]
        date_cible = datetime.strptime(date_examen, "%Y-%m-%d")
        aujourdhui = datetime.now()
        jours_restants = (date_cible - aujourdhui).days
        
        if jours_restants <= 0:
            return "La date d'examen est déjà passée ou c'est aujourd'hui."
        
        if not liste_sujets:
            return "Aucun sujet fourni."
        
        planning = {}
        for i in range(jours_restants):
            jour = aujourdhui.replace(hour=0, minute=0, second=0, microsecond=0)
            jour = jour.fromtimestamp(jour.timestamp() + i * 86400)
            sujet_du_jour = liste_sujets[i % len(liste_sujets)]
            planning[jour.strftime("%Y-%m-%d")] = sujet_du_jour
        
        texte_planning = f"Planning de révision jusqu'au {date_examen} ({jours_restants} jours) :\n"
        for jour, sujet in planning.items():
            texte_planning += f"- {jour} : {sujet}\n"
        
        with open("planning_revisions.txt", "w", encoding="utf-8") as f:
            f.write(texte_planning)
        
        return texte_planning
    except ValueError:
        return "Format de date incorrect. Utilise le format AAAA-MM-JJ (ex: 2026-11-15)."
    except Exception as e:
        return f"Erreur : {e}"
def generer_qcm(texte: str, nombre_questions: int = 5) -> str:
    try:
        client = Groq(api_key=os.getenv("GROQ_API_KEY"))
        
        prompt = f"""À partir du texte suivant, crée {nombre_questions} questions à choix multiples (QCM) 
pour réviser. Chaque question doit avoir 4 propositions (A, B, C, D) et indiquer la bonne réponse.

Texte :
{texte}

Format de réponse :
Question 1 : ...
A) ...
B) ...
C) ...
D) ...
Bonne réponse : ...
"""
        
        reponse = client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[{"role": "user", "content": prompt}]
        )
        
        resultat = reponse.choices[0].message.content
        
        with open("qcm_genere.txt", "w", encoding="utf-8") as f:
            f.write(resultat)
        
        return resultat
    except Exception as e:
        return f"Erreur : {e}"
def charger_memoire_generale() -> str:
    try:
        with open("memoire.txt", "r", encoding="utf-8") as f:
            return f.read().strip()
    except FileNotFoundError:
        return ""

def enrichir_avec_contexte(instruction: str) -> str:
    contexte = charger_memoire_generale()
    if contexte:
        return f"Contexte sur Monsieur Farnèse :\n{contexte}\n\nDemande : {instruction}"
    return instruction
def rechercher_avec_alpha(question: str) -> str:
    from alpha_agent import collecter_informations
    return collecter_informations(enrichir_avec_contexte(question))

def gerer_fichiers_avec_beta(instruction: str) -> str:
    from beta_agent import gerer_fichiers
    return gerer_fichiers(enrichir_avec_contexte(instruction))

def resoudre_avec_grio(instruction: str) -> str:
    from grio_agent import resoudre_mathematiques
    return resoudre_mathematiques(enrichir_avec_contexte(instruction))

def gerer_revisions_avec_delta(instruction: str) -> str:
    from delta_agent import gerer_revisions
    return gerer_revisions(enrichir_avec_contexte(instruction))

def gerer_agenda_temps_avec_gamma(instruction: str) -> str:
    from gamma_agent import gerer_agenda_temps
    return gerer_agenda_temps(enrichir_avec_contexte(instruction))

def gerer_documents_avec_lambda(instruction: str) -> str:
    from lambda_agent import gerer_documents
    return gerer_documents(enrichir_avec_contexte(instruction))
OUTILS_DISPONIBLES ={
    "lire_fichier": lire_fichier,
    "ecrire_fichier": ecrire_fichier,
    "lister_fichiers": lister_fichiers,
    "calculer": calculer,
    "ajouter_a_memoire": ajouter_a_memoire,
    "rechercher_web": rechercher_web,
    "obtenir_date_heure": obtenir_date_heure,
    "creer_flashcard": creer_flashcard,
    "convertir_unite": convertir_unite,
    "compter_mots": compter_mots,
    "obtenir_meteo": obtenir_meteo,
    "envoyer_email": envoyer_email,
    "resumer_memoire": resumer_memoire,
    "creer_skill": creer_skill,
    "lire_pdf": lire_pdf,
    "suivre_progression": suivre_progression,
    "voir_progression": voir_progression,
    "exporter_donnees": exporter_donnees,
    "planifier_revisions": planifier_revisions,
    "generer_qcm": generer_qcm,
    "recherche_avec_alpha":rechercher_avec_alpha,
    "gerer_fichiers_avec_beta": gerer_fichiers_avec_beta,
    "resoudre_avec_grio": resoudre_avec_grio,
    "gerer_revisions_avec_delta": gerer_revisions_avec_delta,
    "gerer_agenda_temps_avec_gamma": gerer_agenda_temps_avec_gamma,
    "lire_docx": lire_docx,
    "ecrire_docx": ecrire_docx,
    "lire_excel": lire_excel,
    "ecrire_excel": ecrire_excel,
    "lire_pptx": lire_pptx,
    "creer_pptx": creer_pptx,
    "gerer_documents_avec_lambda": gerer_documents_avec_lambda,
}
